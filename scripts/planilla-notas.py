#!/usr/bin/env python3
"""Completa en la planilla de fórmulas las columnas de ayuda «Árabe de tus notas» y «Transliteración de tus notas».

Uso: python3 scripts/planilla-notas.py <planilla.xlsx> <carpeta con los «Texto del…»> [--comprobar]

Lee las notas de estudio del usuario (Obsidian, vault Islam: 08 Práctica/Rezos/Texto del fajr.md, …) y copia,
tal cual, la línea árabe y la transliteración de cada fórmula de «Mi práctica» en dos columnas al final de la
hoja «Fórmulas». Son una ayuda para quien revisa: las columnas amarillas (árabe, transliteración y traducción
revisadas, «Revisó») no se tocan, y nada de esto llega a formulas.json si una persona no lo copia y firma.
No transcribe, corrige ni completa texto: si una línea no está literalmente en la nota, falla.

Donde la nota junta dos fórmulas en una línea, se separa por un límite exacto del texto:
el primer punto del tashahhud («Al-ḥamdu lillāh. Ash-hadu…») y «wa raḥmatu» en el saludo final.
Necesita openpyxl (python3 -m pip install openpyxl, o en un venv). Con --comprobar solo informa.
"""
from pathlib import Path
import re
import sys

try:
    import openpyxl
except ImportError:
    raise SystemExit('Falta openpyxl: python3 -m pip install openpyxl (o instalarlo en un venv).')
from copy import copy

NOTAS = {'fajr': 'Texto del fajr.md', 'dhuhr': 'Texto del dhuhr.md', 'asr': 'Texto del asr.md',
         'maghrib': 'Texto del maghrib.md', 'isha': 'Texto de la isha.md'}
ARABE = re.compile(r'[؀-ۿ]')
COL_AR, COL_TR = 'Árabe de tus notas (copiar si está bien)', 'Transliteración de tus notas'


def pares(texto):
    """(sección, árabe, transliteración) de cada línea árabe seguida de su transliteración en cursiva."""
    lineas = [l.strip() for l in texto.split('\n')]
    seccion = ''
    for i, l in enumerate(lineas):
        if l.startswith('## '):
            seccion = l[3:].lower()
        elif ARABE.search(l) and not l.startswith('Tasbih'):
            ar = re.sub(r'\s*(×\s*\d+)?\s*(—.*)?$', '', l)
            j = i + 1
            while j < len(lineas) and not lineas[j]:
                j += 1
            m = re.match(r'^\*([^*]+)\*', lineas[j]) if j < len(lineas) else None
            if m:
                yield seccion, ar, m.group(1)


def tasbih(texto):
    """Las tres frases del tasbīḥ de Fátima: la línea «Tasbih de Fátima: 34 × … · 33 × …» y su transliteración."""
    lineas = [l.strip() for l in texto.split('\n')]
    i = next(n for n, l in enumerate(lineas) if l.startswith('Tasbih de Fátima:'))
    ar = [re.sub(r'^\d+\s*×\s*', '', p.strip()) for p in lineas[i].split(':', 1)[1].split('·')]
    tr = [re.sub(r'^\d+\s*×\s*', '', p.strip()) for p in lineas[i + 1].strip('*').split('·')]
    return list(zip(ar, tr))


def partir(par, sep_ar, sep_tr, conservar):
    """Separa una línea en dos fórmulas; `conservar` deja el separador al comienzo de la segunda."""
    (a1, a2), (t1, t2) = par[0].split(sep_ar, 1), par[1].split(sep_tr, 1)
    if conservar:
        a2, t2 = sep_ar.strip() + a2, sep_tr.strip() + t2
    return (a1.strip(), t1.strip()), (a2.strip(), t2.strip())


def formulas(carpeta):
    textos = {r: (carpeta / n).read_text(encoding='utf-8') for r, n in NOTAS.items()}
    lista = {r: list(pares(t)) for r, t in textos.items()}
    def buscar(rezo, prefijo, seccion=''):
        return next((ar, tr) for s, ar, tr in lista[rezo] if tr.startswith(prefijo) and seccion in s)
    fajr = 'fajr'
    tash = buscar(fajr, 'Al-ḥamdu lillāh. Ash-hadu')
    salam = buscar(fajr, 'As-salāmu ʿalaykum wa')
    hamd, j_tash = partir(tash, '. ', '. ', False)
    j_salam, agregado = partir(salam, ' وَرَحْمَةُ', ' wa raḥmatu', True)
    f = {
        'j_takbir': buscar(fajr, 'Allāhu akbar', 'niyya'), 'j_ruku': buscar(fajr, 'Subḥāna rabbiya l-ʿaẓīm'),
        'j_samia': buscar(fajr, 'Samiʿa llāhu'), 'j_suyud': buscar(fajr, 'Subḥāna rabbiya l-aʿlā'),
        'j_tasbihat': buscar('dhuhr', 'Subḥāna llāhi wa l-ḥamdu'), 'j_tashahhud': j_tash,
        'j_salam_nabi': buscar(fajr, 'As-salāmu ʿalayka'), 'j_salam_alayna': buscar(fajr, 'As-salāmu ʿalaynā'),
        'j_salam': j_salam, 'j_salam_agregado': agregado,
        'adh_tawhid': buscar(fajr, 'Ash-hadu an lā', 'adhan'), 'adh_risala': buscar(fajr, 'Ash-hadu anna Muḥammadan', 'adhan'),
        'adh_wilaya': buscar(fajr, 'Ash-hadu anna ʿAliyyan', 'adhan'), 'adh_salat': buscar(fajr, 'Ḥayya ʿalā ṣ-ṣalāh', 'adhan'),
        'adh_falah': buscar(fajr, 'Ḥayya ʿalā l-falāḥ', 'adhan'), 'adh_khayr': buscar(fajr, 'Ḥayya ʿalā khayri', 'adhan'),
        'adh_qad_qamat': buscar(fajr, 'Qad qāmati', 'iqama'), 'adh_tahlil': buscar(fajr, 'Lā ilāha illā llāh', 'adhan'),
        'm_hamd': (next(ar for ar, tr in tasbih(textos[fajr]) if tr == 'al-ḥamdu lillāh'), hamd[1]),
        'm_salawat': buscar(fajr, 'Allāhumma ṣalli ʿalā Muḥammadin wa āli Muḥammad'),
    }
    for rezo in NOTAS:
        f[f'm_niyya_{rezo}'] = buscar(rezo, 'Nawaytu')
    for fid, (ar, tr) in f.items():  # copia literal: cada tramo está en alguna nota
        if not any(ar in t for t in textos.values()) or not any(tr in t for t in textos.values()):
            raise SystemExit(f'{fid}: el texto no aparece literalmente en las notas')
    return f


def main():
    args = [a for a in sys.argv[1:] if a != '--comprobar']
    if len(args) != 2:
        raise SystemExit(__doc__)
    planilla, carpeta = Path(args[0]), Path(args[1])
    f = formulas(carpeta)
    libro = openpyxl.load_workbook(planilla)
    hoja = libro['Fórmulas']
    col = {c.value: c.column for c in hoja[1]}
    ultima = max(col.values())
    modelo_cab = hoja.cell(1, col['Transliteración: borrador (revisar)'])
    for k, nombre in enumerate((COL_AR, COL_TR), 1):
        if nombre not in col:
            col[nombre] = ultima = ultima + 1
            celda = hoja.cell(1, ultima, nombre)
            celda._style = copy(modelo_cab._style)
            hoja.column_dimensions[celda.column_letter].width = 42 if k == 1 else 38
    llenas, sin_fila = 0, set(f)
    for fila in range(2, hoja.max_row + 1):
        fid = hoja.cell(fila, col['Id (para la app)']).value
        modelo = hoja.cell(fila, col['Transliteración: borrador (revisar)'])
        for nombre, valor in ((COL_AR, f.get(fid, (None, None))[0]), (COL_TR, f.get(fid, (None, None))[1])):
            celda = hoja.cell(fila, col[nombre])
            celda._style = copy(modelo._style)
            celda.value = valor
        if fid in f:
            llenas += 1
            sin_fila.discard(fid)
    print(f'{llenas} filas con texto de tus notas.')
    if sin_fila:
        print('Sin fila en la planilla:', ', '.join(sorted(sin_fila)))
    if '--comprobar' not in sys.argv:
        libro.save(planilla)
        print(f'Guardado en {planilla}.')


if __name__ == '__main__':
    main()
